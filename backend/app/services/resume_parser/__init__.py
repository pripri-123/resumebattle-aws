import io
import re
from typing import Dict, Any, List

class ResumeParserService:
    @staticmethod
    def extract_text(file_bytes: bytes, filename: str) -> str:
        ext = filename.split(".")[-1].lower() if "." in filename else ""
        if ext == "pdf":
            return ResumeParserService._parse_pdf(file_bytes)
        elif ext in ["docx", "doc"]:
            return ResumeParserService._parse_docx(file_bytes)
        else:
            # fallback to utf-8 plain text
            try:
                return file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                return file_bytes.decode("latin-1", errors="ignore")

    @staticmethod
    def _parse_pdf(file_bytes: bytes) -> str:
        from pypdf import PdfReader
        text = ""
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        except Exception as e:
            text = f"Warning: PDF parsing fallback due to: {str(e)}"
        return ResumeParserService.sanitize_text(text)

    @staticmethod
    def _parse_docx(file_bytes: bytes) -> str:
        from docx import Document
        text = ""
        try:
            doc = Document(io.BytesIO(file_bytes))
            for p in doc.paragraphs:
                if p.text:
                    text += p.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        text += " | ".join(row_text) + "\n"
        except Exception as e:
            text = f"Warning: DOCX parsing fallback due to: {str(e)}"
        return ResumeParserService.sanitize_text(text)

    @staticmethod
    def sanitize_text(text: str) -> str:
        if not text:
            return ""
        # Remove null characters and non-printable control codes
        cleaned = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
        # Normalize whitespace
        cleaned = re.sub(r'[ \t]+', ' ', cleaned)
        cleaned = re.sub(r'\n{3,}', '\n\n', cleaned)
        return cleaned.strip()

    @staticmethod
    def extract_sections(text: str) -> Dict[str, str]:
        """Heuristically extracts key resume sections."""
        sections = {
            'summary': '',
            'skills': '',
            'experience': '',
            'projects': '',
            'education': '',
            'certifications': ''
        }
        
        headers = {
            'skills': r'(?:technical\s+)?skills|technologies|proficiencies|competencies',
            'experience': r'(?:work\s+|professional\s+)?experience|employment|work\s+history',
            'projects': r'(?:featured\s+|technical\s+)?projects|portfolio',
            'education': r'education|academic\s+background|qualifications',
            'certifications': r'certifications|licenses|awards',
            'summary': r'(?:professional\s+)?summary|profile|about\s+me'
        }
        
        pattern = r'(?m)^(?:[0-9\.\-\*\s]*)((' + '|'.join(headers.values()) + r'))(?:\s*[:\-]|\s*$)'
        matches = list(re.finditer(pattern, text, re.IGNORECASE))
        
        if not matches:
            return {'body': text}
            
        for i, match in enumerate(matches):
            header_text = match.group(1).lower()
            start = match.end()
            end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
            content = text[start:end].strip()
            
            for key, reg in headers.items():
                if re.search(reg, header_text, re.IGNORECASE):
                    sections[key] = content
                    break
                    
        return sections
